import csv
import json
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

def process_custom_csv(input_csv_path):
    if not os.path.exists(input_csv_path):
        print(f"Fehler: Datei {input_csv_path} nicht gefunden.")
        return

    scored_sites = []
    
    with open(input_csv_path, mode="r", encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row in reader:
            try:
                lat = float(row["lat"])
                lon = float(row["lon"])
                alt = float(row["altitude_m"])
                power_dist = float(row["power_dist_m"])
                missing_ops = int(row["missing_operators"])
                access = row["access"]
                name = row["name"]
                site_id = row["id"]
            except KeyError as e:
                print(f"Fehlerhaftes CSV-Format. Fehlende Spalte: {e}")
                return

            # Scoring-Algorithmus
            deficit_pts = 10.0 if missing_ops >= 3 else (7.5 if missing_ops == 2 else (4.0 if missing_ops == 1 else 1.0))
            topo_pts = 9.5 if alt >= 240 else (7.5 if alt >= 180 else 4.0)
            power_pts = 9.5 if power_dist <= 30 else (7.5 if power_dist <= 100 else (5.0 if power_dist <= 300 else 2.0))
            access_pts = 10.0 if access.lower() == "asphaltiert" else (7.5 if "wirtschaftsweg" in access.lower() else 3.0)

            score = round((deficit_pts * 0.35) + (topo_pts * 0.30) + (power_pts * 0.20) + (access_pts * 0.15), 2)
            prio = "P1" if score >= 8.50 else ("P2" if score >= 7.00 else ("P3" if score >= 5.00 else "P4"))

            scored_sites.append({
                "ID": site_id,
                "Name": name,
                "Breitengrad": lat,
                "Laengengrad": lon,
                "Hoehe_m": alt,
                "Strom_Distanz_m": power_dist,
                "Zuwegung": access,
                "Fehlende_Betreiber": missing_ops,
                "Score": score,
                "Prioritaet": prio
            })

    # In DataFrame konvertieren und nach P1 filtern
    df = pd.DataFrame(scored_sites)
    df_p1 = df[df["Prioritaet"] == "P1"]

    print("=" * 80)
    print(f"Verarbeitung abgeschlossen. Gesamt: {len(df)} | P1-Standorte: {len(df_p1)}")
    print("=" * 80)

    # Excel-Export für P1
    if not df_p1.empty:
        wb = Workbook()
        ws = wb.active
        ws.title = "P1 Standorte"

        for row in dataframe_to_rows(df_p1, index=False, header=True):
            ws.append(row)

        header_fill = PatternFill(start_color="1E7E34", end_color="1E7E34", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='D3D3D3'), right=Side(style='thin', color='D3D3D3'),
            top=Side(style='thin', color='D3D3D3'), bottom=Side(style='thin', color='D3D3D3')
        )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output_xlsx = "Batch_P1_Auswertung.xlsx"
        wb.save(output_xlsx)
        print(f"P1-Ergebnisbericht gespeichert unter: {output_xlsx}")

if __name__ == "__main__":
    import sys
    csv_file = sys.argv[1] if len(sys.argv) > 1 else "beispiel_standorte.csv"
    process_custom_csv(csv_file)
