import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

# 1. Datensatz aus Excel laden
df = pd.read_excel("Gemini_Normalized.xlsx")

# 2. Formatierte Ausgabe-Tabelle erzeugen
wb = Workbook()
ws = wb.active
ws.title = "Gemini Incidents"

# DataFrame in Worksheet schreiben
for r in dataframe_to_rows(df, index=False, header=True):
    ws.append(r)

# 3. P1-Vorfälle rot markieren
red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

severity_col = None
for idx, cell in enumerate(ws[1], start=1):
    if cell.value == "severity":
        severity_col = idx
        break

if severity_col:
    for row in ws.iter_rows(min_row=2):
        if str(row[severity_col - 1].value).strip().upper() == "P1":
            for cell in row:
                cell.fill = red_fill

# 4. Aggregierte Zusammenfassung nach Schweregrad
summary = df.groupby("severity").size().reset_index(name="count")

ws_summary = wb.create_sheet("Summary")
ws_summary.append(["severity", "count"])

for r in dataframe_to_rows(summary, index=False, header=False):
    ws_summary.append(r)

# 5. Datei speichern
wb.save("Gemini_Incident_Report.xlsx")
print("Verarbeitung erfolgreich: Gemini_Incident_Report.xlsx wurde erstellt.")
