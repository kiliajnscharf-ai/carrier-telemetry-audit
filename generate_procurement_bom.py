import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import datetime

def generate_bom():
    print("================================================================================")
    print("PROJEKT HAUS IM WIND: B2B-STUECKLISTEN- & BESCHAFFUNGSGENERATOR (PLATZ 1)")
    print("================================================================================")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stueckliste_BOM"

    navy_fill = PatternFill(start_color="0F243E", end_color="0F243E", fill_type="solid")
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=13, bold=True, color="0F243E")
    bold_font = Font(name="Calibri", size=10, bold=True)
    reg_font = Font(name="Calibri", size=10)
    accent_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.merge_cells("A1:G1")
    ws["A1"] = "B2B-STÜCKLISTE & BESCHAFFUNGSMATRIX: WERKSTATT & MONTAGE (CLUSTER 2026)"
    ws["A1"].font = title_font
    ws["A1"].alignment = left

    headers = ["Pos.", "Hersteller", "Art.-Nr. / Typ", "Bezeichnung / Spezifikation", "Menge", "Einzelpreis (netto)", "Gesamt (netto)"]
    ws.append([])
    ws.append(headers)
    for col in range(1, 8):
        c = ws.cell(row=3, column=col)
        c.fill = navy_fill
        c.font = white_bold
        c.alignment = center

    bom_items = [
        (1, "Fischer", "544165", "FIS EM Plus 390 S Verbundmörtel (390 ml)", 36, 42.50),
        (2, "Fischer", "90451", "FIS A M24 x 300 mm Edelstahl A4-70", 72, 18.90),
        (3, "Wera", "05075625001", "Click-Torque D 6 Drehmomentschlüssel (60-300 Nm)", 2, 285.00),
        (4, "Wera", "05075622001", "Click-Torque C 3 Drehmomentschlüssel (40-200 Nm)", 2, 215.00),
        (5, "Wera", "05059030001", "Kraftform Kompakt VDE 16 extra slim Set", 3, 115.00),
        (6, "Knipex", "97 52 65", "Vierkant-Crimpzange (10-95 mm²)", 2, 340.00),
        (7, "Knipex", "95 16 200", "Kabelschere mit Doppelschneide VDE 1000V", 4, 62.50),
        (8, "Knipex", "87 01 300", "Cobra Hightech-Wasserpumpenzange 300 mm", 4, 38.90),
        (9, "Otto-Chemie", "S110-00", "OTTOSEAL S 110 Premium-Silikon Transparent (310 ml)", 48, 8.75),
        (10, "Otto-Chemie", "S70-00", "OTTOSEAL S 70 Naturstein-Silikon Schiefergrau", 24, 11.20),
        (11, "Otto-Chemie", "FS-00", "OTTO Firestop Silikon Brandschutz EI 120", 18, 16.50),
        (12, "Victron/Custom", "LFP-48-6000", "LiFePO4 48V / 288 kWh Block Cluster-Satz", 3, 14500.00)
    ]

    total_sum = 0.0
    for row_idx, item in enumerate(bom_items, start=4):
        total_row = item[4] * item[5]
        total_sum += total_row
        ws.cell(row=row_idx, column=1, value=item[0]).alignment = center
        ws.cell(row=row_idx, column=2, value=item[1]).alignment = left
        ws.cell(row=row_idx, column=3, value=item[2]).alignment = center
        ws.cell(row=row_idx, column=4, value=item[3]).alignment = left
        ws.cell(row=row_idx, column=5, value=item[4]).alignment = center
        
        c_unit = ws.cell(row=row_idx, column=6, value=item[5])
        c_unit.number_format = '#,##0.00 €'
        c_unit.alignment = right
        
        c_tot = ws.cell(row=row_idx, column=7, value=total_row)
        c_tot.number_format = '#,##0.00 €'
        c_tot.alignment = right

        for col in range(1, 8):
            ws.cell(row=row_idx, column=col).border = thin
            ws.cell(row=row_idx, column=col).font = reg_font

    # Gesamtzeile
    last_row = len(bom_items) + 4
    ws.cell(row=last_row, column=1, value="SUMME").alignment = center
    ws.merge_cells(f"A{last_row}:F{last_row}")
    sum_cell = ws.cell(row=last_row, column=7, value=total_sum)
    sum_cell.number_format = '#,##0.00 €'
    sum_cell.alignment = right

    for col in range(1, 8):
        c = ws.cell(row=last_row, column=col)
        c.font = bold_font
        c.fill = accent_fill
        c.border = thin

    col_widths = {"A": 8, "B": 16, "C": 18, "D": 48, "E": 10, "F": 20, "G": 22}
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    excel_filename = "B2B_PROCUREMENT_BOM.xlsx"
    wb.save(excel_filename)
    print(f"BOM-Excel erfolgreich gespeichert: {excel_filename}")

    txt_content = f"""================================================================================
B2B-BESCHAFFUNGS- UND STUECKLISTEN-DOSSIER (PROJEKT HAUS IM WIND)
================================================================================
Datum:             {datetime.datetime.now().strftime('%d.%m.%Y %H:%M:%S')}
Projektstandard:   Tier-1 Global Reference Architecture (Platz 1)
Ziel-Cluster:      LOC-30, LOC-31, LOC-32

KERNKOMPONENTEN & INVESTITION:
- Fischer Befestigungstechnik (Mörtelextrakt & A4-Ankerstangen):  2.890,80 EUR
- Wera Drehmoment- & VDE-Präzisionswerkzeuge:                    1.345,00 EUR
- Knipex Crimp-, Schneid- und Installationswerkzeuge:            1.085,60 EUR
- Otto-Chemie Hochleistungs- und Brandschutz-Dichtstoffe:        985,80 EUR
- 72h-LiFePO4-Hochleistungs-USV-Einheiten (288 kWh x 3):        43.500,00 EUR
--------------------------------------------------------------------------------
BESCHAFFUNGSVOLUMEN GESAMT (NETTO):                            49.807,20 EUR
================================================================================
STATUS: STUECKLISTE BEREIT ZUR BESTELLUNG BEI ZUGELASSENEN FACHHAENDLERN.
================================================================================
"""
    with open("B2B_PROCUREMENT_BOM.txt", "w", encoding="utf-8") as f:
        f.write(txt_content)
    print(txt_content)

if __name__ == '__main__':
    generate_bom()
