import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

# 1. CSV einlesen
df = pd.read_csv("p1_standorte_einreichung.csv", sep=";")

# 2. Excel-Workbook initialisieren
wb = Workbook()
ws = wb.active
ws.title = "P1 Standorte"

# 3. Daten einfügen
for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
    ws.append(row)

# 4. Formatierung: Grüner Header, Rahmen, zentrierte Werte
header_fill = PatternFill(start_color="1E7E34", end_color="1E7E34", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

# Header formatieren
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Datenzeilen formatieren
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
    for cell in row:
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

# 5. Spaltenbreiten automatisch anpassen
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# 6. Datei speichern
wb.save("P1_Mobilfunk_Standorte.xlsx")
print("Excel-Bericht erfolgreich generiert: P1_Mobilfunk_Standorte.xlsx")
