import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

# -----------------------------
# 1. Normalisierter Gemini-Datensatz
# -----------------------------
data = [
    {"incident_id": 1, "severity": "P1", "log_id": "LOG-01", "description": "/dev/sda I/O-Lesefehler", "validated": "OK"},
    {"incident_id": 4, "severity": "P1", "log_id": "LOG-04", "description": "smartd unlesbare Sektoren", "validated": "OK"},
    {"incident_id": 5, "severity": "P2", "log_id": "LOG-05", "description": "DHCP-Pool-Erschöpfung 92%", "validated": "OK"},
    {"incident_id": 6, "severity": "P2", "log_id": "LOG-06", "description": "VPN-Handshake-Timeout", "validated": "OK"},
    {"incident_id": 2, "severity": "P3", "log_id": "LOG-02", "description": "DNS-Failover auf Public DNS", "validated": "OK"},
    {"incident_id": 3, "severity": "P4", "log_id": "LOG-03", "description": "Nginx-API-Zugriff 200 OK", "validated": "OK"},
]

df = pd.DataFrame(data)

# -----------------------------
# 2. Excel-Workbook erzeugen
# -----------------------------
wb = Workbook()

# -----------------------------
# 3. Haupttabelle
# -----------------------------
ws = wb.active
ws.title = "Gemini Incidents"

for row in dataframe_to_rows(df, index=False, header=True):
    ws.append(row)

# -----------------------------
# 4. P1 rot markieren
# -----------------------------
red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

severity_col = None
for idx, cell in enumerate(ws[1], start=1):
    if cell.value == "severity":
        severity_col = idx
        break

if severity_col:
    for row in ws.iter_rows(min_row=2):
        if str(row[severity_col - 1].value).strip().upper() == "P1":
            for cell in row:
                cell.fill = red_fill

# -----------------------------
# 5. Zusammenfassung nach Schweregrad
# -----------------------------
summary = df.groupby("severity").size().reset_index(name="count")

ws_summary = wb.create_sheet("Summary")
ws_summary.append(["severity", "count"])

for row in dataframe_to_rows(summary, index=False, header=False):
    ws_summary.append(row)

# -----------------------------
# 6. Speichern
# -----------------------------
wb.save("Gemini_Incident_Report.xlsx")
print("Version 2 erfolgreich ausgeführt: Gemini_Incident_Report.xlsx aktualisiert.")
